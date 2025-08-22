from docx import Document
import random

# Generate 50 students with random marks and grades
names = [f"Student{i+1}" for i in range(50)]
marks = [random.randint(50, 100) for _ in range(50)]
grades = []
for m in marks:
    if m >= 90:
        grades.append("A")
    elif m >= 80:
        grades.append("B")
    elif m >= 70:
        grades.append("C")
    elif m >= 60:
        grades.append("D")
    else:
        grades.append("F")

students = list(zip(names, marks, grades))

# Create Word document
doc = Document()
doc.add_heading('Student Marks and Grades', 0)

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Student'
hdr_cells[1].text = 'Marks'
hdr_cells[2].text = 'Grade'

for name, mark, grade in students:
    row_cells = table.add_row().cells
    row_cells[0].text = name
    row_cells[1].text = str(mark)
    row_cells[2].text = grade

doc.save('students_info.docx')
print("Word file 'students_info.docx' created successfully.")