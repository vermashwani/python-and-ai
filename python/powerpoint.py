from pptx import Presentation
from pptx.util import Inches
import matplotlib.pyplot as plt

# Student data
students = [
    ("John", 85, "A"),
    ("Sara", 78, "B"),
    ("Mike", 92, "A"),
    ("Lily", 67, "C"),
]

# Create a bar chart and save as image
names = [s[0] for s in students]
marks = [s[1] for s in students]

plt.bar(names, marks, color='skyblue')
plt.title("Student Marks")
plt.xlabel("Student")
plt.ylabel("Marks")
plt.savefig("chart.png")
plt.close()

# Create PowerPoint presentation
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])

# Add table to slide
rows = len(students) + 1
cols = 3
table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(0.5), Inches(4), Inches(1.5)).table

# Set column headers
table.cell(0, 0).text = "Student"
table.cell(0, 1).text = "Marks"
table.cell(0, 2).text = "Grade"

# Fill table data
for i, (name, marks, grade) in enumerate(students, start=1):
    table.cell(i, 0).text = name
    table.cell(i, 1).text = str(marks)
    table.cell(i, 2).text = grade

# Add chart image to slide
slide.shapes.add_picture("chart.png", Inches(5), Inches(0.5), Inches(4), Inches(3))

# Save PowerPoint file
prs.save("sample_presentation.pptx")
print("PowerPoint file 'sample_presentation.pptx' created successfully.")